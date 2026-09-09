#!/usr/bin/env python3
"""Rebuild the spreadsheet from the CSV after find_phones.py has updated it.

    python3 csv_to_xlsx.py data/florida-marine-businesses.csv
"""
import csv
import os
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

src = sys.argv[1] if len(sys.argv) > 1 else 'data/florida-marine-businesses.csv'
out = sys.argv[2] if len(sys.argv) > 2 else os.path.splitext(src)[0] + '.xlsx'

rows = list(csv.DictReader(open(src, encoding='utf-8')))
cols = list(rows[0].keys())

WIDTH = {'Company': 42, 'Category': 18, 'Phone': 16, 'Phone source': 17,
         'Email': 32, 'Website': 40, 'Address': 46, 'City': 20, 'County': 16,
         'Rating': 8, 'Reviews': 9, 'Status': 18, 'Map': 30, 'OSM': 30}

wb = Workbook()
ws = wb.active
ws.title = 'FL Marine Businesses'
ws.append(cols)
for c in ws[1]:
    c.font = Font(bold=True, color='FFFFFF')
    c.fill = PatternFill('solid', fgColor='11314F')
    c.alignment = Alignment(vertical='center')
ws.freeze_panes = 'A2'
for r in rows:
    ws.append([r[c] for c in cols])
for i, c in enumerate(cols, 1):
    ws.column_dimensions[get_column_letter(i)].width = WIDTH.get(c, 20)
ws.auto_filter.ref = f'A1:{get_column_letter(len(cols))}{len(rows)+1}'
wb.save(out)

print(f'{out}  ({len(rows)} rows, {sum(1 for r in rows if r.get("Phone"))} with a phone)')
